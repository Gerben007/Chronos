"""Excel → entries.json roundtrip.

Parents author timeline entries by editing a single .xlsx in the Chronos
folder of Nextcloud. The ingest loop watches that file, parses it, and
overwrites db/seed/entries.json. The site rebuild then bakes the new
list into the static dist.

Schema (one row per entry, header on row 1):

  slug                  required, lowercase ASCII slug, unique
  type                  required, one of: person | event | period | prophecy
  lane_slug             required, must match a lane in db/seed/lanes.json
  start_year            required, signed int (BC = negative)
  end_year              optional, signed int
  importance            required, 1-5
  title_en              required
  title_af              optional (falls back to title_en if blank)
  display_dates_en      required
  display_dates_af      optional
  summary_en            required
  summary_af            optional
  scripture_refs        optional, semicolon-separated list
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any


COLUMNS = [
    "slug",
    "type",
    "lane_slug",
    "start_year",
    "end_year",
    "importance",
    "title_en",
    "title_af",
    "display_dates_en",
    "display_dates_af",
    "summary_en",
    "summary_af",
    "scripture_refs",
]

VALID_TYPES = {"person", "event", "period", "prophecy"}
SLUG_RE = re.compile(r"^[a-z0-9][a-z0-9-]*$")


@dataclass
class ParseResult:
    entries: list[dict[str, Any]] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)        # row-level
    seen_slugs: set[str] = field(default_factory=set)


# ── Parse ──────────────────────────────────────────────────────────────


def parse_xlsx(path: Path, *, valid_lanes: set[str]) -> ParseResult:
    """Reads the first worksheet. Header row drives column mapping.

    Returns a ParseResult — caller decides whether to abort on errors
    or accept partial output. We never raise on a bad row; we collect
    the error and move on so one typo doesn't take down the timeline.
    """
    from openpyxl import load_workbook    # imported lazily for test reach

    result = ParseResult()
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    if ws is None or ws.max_row < 2:
        result.errors.append("worksheet empty")
        return result

    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter)
    header_map = _build_header_map(header)
    missing = [c for c in COLUMNS if c.endswith("en") or c in ("slug","type","lane_slug","start_year","importance")]
    missing = [c for c in missing if c not in header_map]
    if missing:
        result.errors.append(f"missing required columns: {', '.join(missing)}")
        return result

    for idx, row in enumerate(rows_iter, start=2):
        if all(v is None or str(v).strip() == "" for v in row):
            continue       # blank row, skip
        try:
            entry = _row_to_entry(row, header_map, valid_lanes=valid_lanes)
        except ValueError as e:
            result.errors.append(f"row {idx}: {e}")
            continue
        if entry["slug"] in result.seen_slugs:
            result.errors.append(f"row {idx}: duplicate slug {entry['slug']}")
            continue
        result.seen_slugs.add(entry["slug"])
        result.entries.append(entry)

    result.entries.sort(key=lambda e: e["start_year"])
    return result


def _build_header_map(header: tuple) -> dict[str, int]:
    out: dict[str, int] = {}
    for i, name in enumerate(header):
        if name is None:
            continue
        key = str(name).strip().lower().replace(" ", "_")
        if key in COLUMNS:
            out[key] = i
    return out


def _row_to_entry(row: tuple, headers: dict[str, int], *, valid_lanes: set[str]) -> dict[str, Any]:
    def get(col: str) -> Any:
        i = headers.get(col)
        if i is None:
            return None
        v = row[i]
        if isinstance(v, str):
            v = v.strip()
            return v or None
        return v

    slug = get("slug")
    if not isinstance(slug, str) or not SLUG_RE.match(slug):
        raise ValueError(f"slug invalid: {slug!r}")

    type_ = get("type")
    if type_ not in VALID_TYPES:
        raise ValueError(f"type must be one of {sorted(VALID_TYPES)}, got {type_!r}")

    lane_slug = get("lane_slug")
    if lane_slug not in valid_lanes:
        raise ValueError(f"lane_slug {lane_slug!r} not in {sorted(valid_lanes)}")

    start_year = get("start_year")
    if not isinstance(start_year, int):
        raise ValueError(f"start_year must be int, got {start_year!r}")
    end_year = get("end_year")
    if end_year is not None and not isinstance(end_year, int):
        raise ValueError(f"end_year must be int or blank, got {end_year!r}")

    importance = get("importance")
    if not isinstance(importance, int) or not 1 <= importance <= 5:
        raise ValueError(f"importance must be 1-5, got {importance!r}")

    title_en = get("title_en")
    if not isinstance(title_en, str) or not title_en:
        raise ValueError("title_en required")
    display_dates_en = get("display_dates_en")
    if not isinstance(display_dates_en, str) or not display_dates_en:
        raise ValueError("display_dates_en required")
    summary_en = get("summary_en")
    if not isinstance(summary_en, str) or not summary_en:
        raise ValueError("summary_en required")

    title_af = get("title_af") or title_en
    display_dates_af = get("display_dates_af") or display_dates_en
    summary_af = get("summary_af") or summary_en

    refs_raw = get("scripture_refs")
    scripture_refs: list[str] = []
    if isinstance(refs_raw, str) and refs_raw:
        scripture_refs = [r.strip() for r in refs_raw.split(";") if r.strip()]

    entry: dict[str, Any] = {
        "slug": slug,
        "type": type_,
        "lane_slug": lane_slug,
        "start_year": start_year,
        "importance": importance,
        "translations": {
            "en": {
                "title": title_en,
                "display_dates": display_dates_en,
                "summary": summary_en,
            },
            "af": {
                "title": title_af,
                "display_dates": display_dates_af,
                "summary": summary_af,
            },
        },
    }
    if end_year is not None:
        entry["end_year"] = end_year
    if scripture_refs:
        entry["scripture_refs"] = scripture_refs
    return entry


# ── Export (entries.json → .xlsx template) ────────────────────────────


def export_xlsx(entries: list[dict[str, Any]], dest: Path) -> None:
    """Writes a workbook with the canonical schema, prefilled from entries.

    Used to bootstrap parents: dump the existing seed to .xlsx, drop it
    in Nextcloud, and they have a working template they can extend.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "entries"

    # Header row
    for col_idx, name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DCD2BE")
    ws.freeze_panes = "A2"

    # Data rows
    for row_idx, e in enumerate(entries, start=2):
        en = e.get("translations", {}).get("en", {})
        af = e.get("translations", {}).get("af", {})
        values = {
            "slug":             e.get("slug"),
            "type":             e.get("type"),
            "lane_slug":        e.get("lane_slug"),
            "start_year":       e.get("start_year"),
            "end_year":         e.get("end_year"),
            "importance":       e.get("importance"),
            "title_en":         en.get("title"),
            "title_af":         af.get("title"),
            "display_dates_en": en.get("display_dates"),
            "display_dates_af": af.get("display_dates"),
            "summary_en":       en.get("summary"),
            "summary_af":       af.get("summary"),
            "scripture_refs":   "; ".join(e.get("scripture_refs", []) or []),
        }
        for col_idx, name in enumerate(COLUMNS, start=1):
            ws.cell(row=row_idx, column=col_idx, value=values[name])

    # Generous widths
    widths = {"slug": 22, "type": 12, "lane_slug": 18, "start_year": 12, "end_year": 12,
              "importance": 12, "title_en": 28, "title_af": 28,
              "display_dates_en": 22, "display_dates_af": 22,
              "summary_en": 60, "summary_af": 60, "scripture_refs": 40}
    for col_idx, name in enumerate(COLUMNS, start=1):
        ws.column_dimensions[chr(64 + col_idx) if col_idx < 27 else "A" + chr(38 + col_idx)].width = widths[name]

    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
